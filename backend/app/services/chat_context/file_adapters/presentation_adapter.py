import os
import logging
from typing import Dict, List, Optional, Any
from app.services.chat_context.contracts import (
    FileCapability,
    FileProcessingState,
    DocumentChunk,
    FileAnalysisResult
)
from app.services.chat_context.file_adapters.base import BaseFileAdapter

logger = logging.getLogger("hsbot.chat_context.presentation")


class PresentationAdapter(BaseFileAdapter):
    capability = FileCapability.PRESENTATION
    supported_extensions = [".pptx", ".ppt", ".odp", ".potx"]
    supported_mimes = [
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.ms-powerpoint",
        "application/vnd.oasis.opendocument.presentation"
    ]

    def _estimate_tokens(self, text: str) -> int:
        return max(1, len(text) // 4)

    async def extract(self, file_path: str, filename: str, file_id: str) -> FileAnalysisResult:
        if not os.path.exists(file_path):
            return FileAnalysisResult(
                file_id=file_id,
                filename=filename,
                capability=self.capability,
                state=FileProcessingState.FAILED,
                error=f"File not found: {file_path}"
            )

        ext = os.path.splitext(filename)[1].lower()
        if ext in [".ppt", ".odp"]:
            return FileAnalysisResult(
                file_id=file_id,
                filename=filename,
                capability=self.capability,
                state=FileProcessingState.PARTIALLY_SUPPORTED,
                summary=f"Legacy or OpenDocument presentation ({ext}). For complete slide structure, please save as .pptx.",
                is_partially_supported=True
            )

        chunks: List[DocumentChunk] = []
        citations: List[str] = []
        slide_summaries: List[str] = []

        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            total_slides = len(prs.slides)

            for slide_idx, slide in enumerate(prs.slides):
                slide_num = slide_idx + 1
                slide_title = ""
                slide_text_lines = []

                # Extract slide title if available
                if slide.shapes.title and slide.shapes.title.text:
                    slide_title = slide.shapes.title.text.strip()

                for shape in slide.shapes:
                    if shape == slide.shapes.title:
                        continue
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            t = paragraph.text.strip()
                            if t:
                                slide_text_lines.append(t)
                    elif shape.has_table:
                        table = shape.table
                        t_rows = []
                        for row in table.rows:
                            t_rows.append(" | ".join(cell.text.strip() for cell in row.cells))
                        if t_rows:
                            slide_text_lines.append("\n".join(t_rows))

                # Speaker notes
                notes_text = ""
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    nt = slide.notes_slide.notes_text_frame.text.strip()
                    if nt:
                        notes_text = f"Speaker Notes: {nt}"

                title_header = f"Slide {slide_num}: {slide_title}" if slide_title else f"Slide {slide_num}"
                body_content = "\n".join(slide_text_lines)
                full_slide_content = f"### {title_header}\n{body_content}"
                if notes_text:
                    full_slide_content += f"\n\n{notes_text}"

                if not slide_text_lines and not slide_title and not notes_text:
                    continue

                chunk_id = f"{file_id}_slide_{slide_num}"
                location = f"Slide {slide_num} of {filename}"
                token_count = self._estimate_tokens(full_slide_content)

                chunk = DocumentChunk(
                    chunk_id=chunk_id,
                    file_id=file_id,
                    source=filename,
                    location=location,
                    section=title_header,
                    slide=slide_num,
                    content=full_slide_content,
                    token_count=token_count,
                    metadata={"slide_num": slide_num, "title": slide_title}
                )
                chunks.append(chunk)
                citations.append(location)
                slide_summaries.append(f"Slide {slide_num}: {slide_title or '(No Title)'}")

            total_tokens = sum(c.token_count for c in chunks)
            summary = (
                f"Presentation: '{filename}' ({total_slides} slides, {len(chunks)} text slides extracted, ~{total_tokens} tokens).\n"
                + "\n".join(f"- {s}" for s in slide_summaries[:10])
            )

            return FileAnalysisResult(
                file_id=file_id,
                filename=filename,
                capability=self.capability,
                state=FileProcessingState.READY,
                summary=summary,
                chunks=chunks,
                total_tokens=total_tokens,
                total_pages=total_slides,
                metadata={"total_slides": total_slides, "extracted_slides": len(chunks)},
                citations=citations
            )

        except Exception as e:
            logger.error(f"Error parsing PPTX presentation: {e}")
            return FileAnalysisResult(
                file_id=file_id,
                filename=filename,
                capability=self.capability,
                state=FileProcessingState.FAILED,
                error=f"Error reading presentation: {str(e)}"
            )
