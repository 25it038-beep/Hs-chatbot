import os
from typing import List, Dict, Optional, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def generate_docx(
    title: str,
    sections: List[Dict[str, Any]],
    output_path: str,
    author: Optional[str] = None,
    subtitle: Optional[str] = None,
) -> str:
    """Generates a styled Word document."""
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = Document()

    # Document Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = title_para.add_run(title)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    # Subtitle / Author
    if subtitle or author:
        meta_para = doc.add_paragraph()
        meta_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        meta_text = " • ".join(filter(None, [subtitle, f"Author: {author}" if author else None]))
        run_meta = meta_para.add_run(meta_text)
        run_meta.font.name = "Calibri"
        run_meta.font.size = Pt(11)
        run_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()  # spacer

    # Sections
    for sec in sections:
        heading = sec.get("heading", "")
        content = sec.get("content", "")
        items = sec.get("items", [])
        table_data = sec.get("table")

        if heading:
            h_para = doc.add_heading(level=1)
            h_run = h_para.add_run(heading)
            h_run.font.name = "Calibri"
            h_run.font.size = Pt(16)
            h_run.font.bold = True
            h_run.font.color.rgb = RGBColor(30, 41, 59)

        if content:
            for para in content.split("\n\n"):
                clean_p = para.strip()
                if not clean_p:
                    continue
                if clean_p.startswith("- ") or clean_p.startswith("* ") or clean_p.startswith("• "):
                    for line in clean_p.split("\n"):
                        l = line.strip().lstrip("-*• ").strip()
                        if l:
                            doc.add_paragraph(l, style='List Bullet')
                else:
                    p = doc.add_paragraph(clean_p)
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_after = Pt(6)

        if items:
            for itm in items:
                doc.add_paragraph(str(itm), style='List Bullet')

        if table_data and isinstance(table_data, list) and len(table_data) > 0:
            num_cols = len(table_data[0])
            num_rows = len(table_data)
            table = doc.add_table(rows=num_rows, cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Light Shading Accent 1' if 'Light Shading Accent 1' in [s.name for s in doc.styles] else 'Table Grid'

            for r_idx, row in enumerate(table_data):
                for c_idx, cell_value in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = str(cell_value)
                    if r_idx == 0:
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.bold = True

            doc.add_paragraph()  # spacer

    doc.save(output_path)
    return output_path


def generate_simple_docx(title: str, text: str, output_path: str) -> str:
    """Fallback generator converting markdown/text lines to docx paragraphs."""
    sections = []
    current_sec = {"heading": "Overview", "content": ""}

    for line in text.splitlines():
        trimmed = line.strip()
        if trimmed.startswith("#"):
            if current_sec["content"].strip():
                sections.append(current_sec)
            h = trimmed.lstrip("#").strip()
            current_sec = {"heading": h, "content": ""}
        else:
            current_sec["content"] += line + "\n"

    if current_sec["content"].strip():
        sections.append(current_sec)

    if not sections:
        sections = [{"heading": "Content", "content": text}]

    return generate_docx(title=title, sections=sections, output_path=output_path)
