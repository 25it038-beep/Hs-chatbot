import os
import csv
import logging
from typing import Tuple

logger = logging.getLogger("hsbot.document.validator")


def validate_pdf(path: str) -> bool:
    valid, _ = validate_file_structure(path, "pdf")
    return valid


def validate_docx(path: str) -> bool:
    valid, _ = validate_file_structure(path, "docx")
    return valid


def validate_pptx(path: str) -> bool:
    valid, _ = validate_file_structure(path, "pptx")
    return valid


def validate_xlsx(path: str) -> bool:
    valid, _ = validate_file_structure(path, "xlsx")
    return valid


def validate_csv(path: str) -> bool:
    valid, _ = validate_file_structure(path, "csv")
    return valid


def validate_file_structure(path: str, fmt: str) -> Tuple[bool, str]:
    """Validates that a generated file exists, is non-empty, has the correct extension,

    and possesses valid binary/package structure.
    """
    if not os.path.exists(path):
        return False, f"File does not exist: {path}"

    try:
        file_size = os.path.getsize(path)
    except Exception as e:
        return False, f"Cannot get file size: {e}"

    if file_size <= 0:
        return False, f"File is empty (0 bytes): {path}"

    fmt = fmt.lower().strip(".")
    ext = os.path.splitext(path)[1].lower().strip(".")
    if ext != fmt and not (fmt in ("markdown", "md") and ext in ("md", "markdown")):
        return False, f"Extension mismatch: expected .{fmt}, got .{ext}"

    try:
        with open(path, "rb") as f:
            header = f.read(16)
    except Exception as e:
        return False, f"Failed to read file header: {e}"

    if fmt == "pdf":
        if not header.startswith(b"%PDF-"):
            return False, "Invalid PDF magic header (does not start with %PDF-)"
        return True, "Valid PDF"

    elif fmt == "docx":
        if not header.startswith(b"PK\x03\x04"):
            return False, "Invalid DOCX package (missing ZIP PK header)"
        try:
            from docx import Document
            doc = Document(path)
            # basic check that document loads
            _ = len(doc.paragraphs)
            return True, "Valid DOCX"
        except Exception as e:
            return False, f"Corrupt DOCX document: {e}"

    elif fmt == "pptx":
        if not header.startswith(b"PK\x03\x04"):
            return False, "Invalid PPTX package (missing ZIP PK header)"
        try:
            from pptx import Presentation
            prs = Presentation(path)
            _ = len(prs.slides)
            return True, "Valid PPTX"
        except Exception as e:
            return False, f"Corrupt PPTX presentation: {e}"

    elif fmt == "xlsx":
        if not header.startswith(b"PK\x03\x04"):
            return False, "Invalid XLSX package (missing ZIP PK header)"
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True)
            _ = wb.sheetnames
            wb.close()
            return True, "Valid XLSX"
        except Exception as e:
            return False, f"Corrupt XLSX workbook: {e}"

    elif fmt == "csv":
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                rows = list(reader)
                if not rows:
                    return False, "CSV contains no data rows"
            return True, "Valid CSV"
        except Exception as e:
            return False, f"Invalid CSV: {e}"

    elif fmt in ("md", "markdown", "txt"):
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
                if not content.strip():
                    return False, "File contains only whitespace"
            return True, "Valid text/markdown"
        except Exception as e:
            return False, f"Invalid text file: {e}"

    return True, "File verified"
