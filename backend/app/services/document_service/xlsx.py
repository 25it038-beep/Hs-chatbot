import os
from typing import List, Dict, Union, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_xlsx(title: str, sheets_data: Dict[str, List[List[Any]]], output_path: str) -> str:
    """Generates a styled Excel workbook with one or more sheets."""
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb = Workbook()
    # Remove initial default sheet
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    regular_font = Font(name="Calibri", size=10, color="0F172A")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    for sheet_name, rows in sheets_data.items():
        clean_sheet_name = (sheet_name or "Sheet1")[:31].replace(":", "").replace("/", "").replace("\\", "").replace("?", "").replace("*", "").replace("[", "").replace("]", "")
        ws = wb.create_sheet(title=clean_sheet_name or "Sheet1")

        for row_idx, row in enumerate(rows, 1):
            for col_idx, val in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border

                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = regular_font
                    if row_idx % 2 == 0:
                        cell.fill = zebra_fill
                    # Right align numbers
                    if isinstance(val, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-size columns with padding
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)
    return output_path


def generate_simple_xlsx(sheet_name: str, data: List[List[Any]], output_path: str) -> str:
    """Generates an Excel workbook from a single 2D array."""
    return generate_xlsx(title=sheet_name, sheets_data={sheet_name: data}, output_path=output_path)
