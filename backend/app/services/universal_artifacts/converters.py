import os
import csv
import json
from typing import Dict, Any, Tuple
import openpyxl
import docx
from app.services.universal_artifacts.generators.document_gen import (
    generate_universal_pdf,
    generate_universal_docx,
    generate_universal_markdown,
)
from app.services.universal_artifacts.generators.spreadsheet_gen import (
    generate_universal_xlsx,
    generate_universal_csv,
)


class ArtifactConverter:
    """Safe Format Conversion Engine (§21, §48, §61)."""

    @classmethod
    def can_convert(cls, from_fmt: str, to_fmt: str) -> bool:
        src = from_fmt.lower().lstrip(".")
        tgt = to_fmt.lower().lstrip(".")
        supported_pairs = {
            ("xlsx", "csv"),
            ("csv", "xlsx"),
            ("docx", "pdf"),
            ("md", "pdf"),
            ("txt", "pdf"),
            ("json", "csv"),
            ("csv", "json"),
        }
        return (src, tgt) in supported_pairs

    @classmethod
    def convert(cls, src_path: str, from_fmt: str, to_fmt: str, target_path: str) -> Tuple[bool, str]:
        src = from_fmt.lower().lstrip(".")
        tgt = to_fmt.lower().lstrip(".")

        if not os.path.exists(src_path):
            return False, f"Source file does not exist: {src_path}"

        try:
            # 1. XLSX -> CSV
            if src == "xlsx" and tgt == "csv":
                wb = openpyxl.load_workbook(src_path, data_only=True)
                ws = wb.active
                with open(target_path, "w", newline="", encoding="utf-8") as f:
                    writer = csv.writer(f)
                    for row in ws.iter_rows(values_only=True):
                        writer.writerow([r if r is not None else "" for r in row])
                wb.close()
                return True, "Converted XLSX to CSV successfully"

            # 2. CSV -> XLSX
            elif src == "csv" and tgt == "xlsx":
                with open(src_path, "r", encoding="utf-8", errors="replace") as f:
                    reader = csv.reader(f)
                    rows = list(reader)
                headers = rows[0] if rows else []
                data_rows = rows[1:] if len(rows) > 1 else []
                generate_universal_xlsx("Converted Spreadsheet", {"headers": headers, "rows": data_rows}, target_path)
                return True, "Converted CSV to XLSX successfully"

            # 3. DOCX -> PDF
            elif src in ["docx", "doc"] and tgt == "pdf":
                doc = docx.Document(src_path)
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                sections = [{"heading": "Converted Document", "paragraphs": paragraphs}]
                generate_universal_pdf("Converted Document", {"sections": sections}, target_path)
                return True, "Converted DOCX to PDF successfully"

            # 4. MD -> PDF
            elif src in ["md", "markdown"] and tgt == "pdf":
                with open(src_path, "r", encoding="utf-8") as f:
                    lines = f.readlines()
                paragraphs = [l.strip() for l in lines if l.strip()]
                sections = [{"heading": "Document Content", "paragraphs": paragraphs}]
                generate_universal_pdf("Markdown Export", {"sections": sections}, target_path)
                return True, "Converted MD to PDF successfully"

            # 5. TXT -> PDF
            elif src == "txt" and tgt == "pdf":
                with open(src_path, "r", encoding="utf-8", errors="replace") as f:
                    paragraphs = [l.strip() for l in f.readlines() if l.strip()]
                sections = [{"heading": "Text Export", "paragraphs": paragraphs}]
                generate_universal_pdf("Text Export", {"sections": sections}, target_path)
                return True, "Converted TXT to PDF successfully"

            # 6. JSON -> CSV
            elif src == "json" and tgt == "csv":
                with open(src_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
                    headers = list(data[0].keys())
                    rows = [[item.get(h, "") for h in headers] for item in data]
                    with open(target_path, "w", newline="", encoding="utf-8") as f:
                        writer = csv.writer(f)
                        writer.writerow(headers)
                        writer.writerows(rows)
                    return True, "Converted JSON array to CSV successfully"
                else:
                    return False, "JSON structure is not an array of objects; cannot convert to CSV"

            # 7. CSV -> JSON
            elif src == "csv" and tgt == "json":
                with open(src_path, "r", encoding="utf-8", errors="replace") as f:
                    reader = csv.DictReader(f)
                    records = list(reader)
                with open(target_path, "w", encoding="utf-8") as f:
                    json.dump(records, f, indent=2)
                return True, "Converted CSV to JSON records successfully"

            return False, f"Unsupported conversion from .{src} to .{tgt}"

        except Exception as e:
            return False, f"Conversion failed: {str(e)}"
