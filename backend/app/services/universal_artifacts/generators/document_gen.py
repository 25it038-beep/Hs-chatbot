import os
import re
from typing import Dict, Any, List, Optional
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def generate_universal_pdf(title: str, content: Dict[str, Any], output_path: str):
    """Generates a professional PDF document via ReportLab (§5, §33)."""
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=12,
        fontName='Helvetica-Bold'
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontSize=15,
        leading=19,
        textColor=colors.HexColor('#0f172a'),
        spaceBefore=14,
        spaceAfter=6,
        fontName='Helvetica-Bold'
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=10,
        leading=15,
        textColor=colors.HexColor('#334155'),
        spaceAfter=8,
    )

    story = []
    # Title & Header
    story.append(Paragraph(title, title_style))
    story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor('#3b82f6'), spaceBefore=2, spaceAfter=14))

    # Abstract / Overview
    overview = content.get("overview") or content.get("summary") or content.get("description")
    if overview:
        story.append(Paragraph(overview, body_style))
        story.append(Spacer(1, 10))

    # Sections
    sections = content.get("sections") or []
    if isinstance(sections, list):
        for sec in sections:
            sec_title = sec.get("heading") or sec.get("title") or "Section"
            story.append(Paragraph(sec_title, h2_style))
            
            paragraphs = sec.get("paragraphs") or sec.get("content") or []
            if isinstance(paragraphs, str):
                story.append(Paragraph(paragraphs, body_style))
            elif isinstance(paragraphs, list):
                for p in paragraphs:
                    if isinstance(p, str):
                        story.append(Paragraph(p, body_style))

            # Optional table in section
            table_data = sec.get("table")
            if table_data and isinstance(table_data, list) and len(table_data) > 0:
                t = Table(table_data)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, -1), 9),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
                ]))
                story.append(Spacer(1, 6))
                story.append(t)
                story.append(Spacer(1, 10))

    doc.build(story)


def generate_universal_docx(title: str, content: Dict[str, Any], output_path: str):
    """Generates a structured Word document via python-docx (§5, §32)."""
    doc = docx.Document()

    # Title
    t_para = doc.add_paragraph()
    t_run = t_para.add_run(title)
    t_run.font.size = Pt(24)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(30, 41, 59)
    t_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # Overview
    overview = content.get("overview") or content.get("summary") or content.get("description")
    if overview:
        p = doc.add_paragraph(overview)
        p.runs[0].font.size = Pt(11)
        doc.add_paragraph()

    # Sections
    sections = content.get("sections") or []
    if isinstance(sections, list):
        for sec in sections:
            sec_title = sec.get("heading") or sec.get("title") or "Section"
            h = doc.add_heading(sec_title, level=2)
            h.runs[0].font.color.rgb = RGBColor(15, 23, 42)

            paragraphs = sec.get("paragraphs") or sec.get("content") or []
            if isinstance(paragraphs, str):
                doc.add_paragraph(paragraphs)
            elif isinstance(paragraphs, list):
                for p_text in paragraphs:
                    if isinstance(p_text, str):
                        doc.add_paragraph(p_text)

            # Table
            table_data = sec.get("table")
            if table_data and isinstance(table_data, list) and len(table_data) > 0:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.style = 'Table Grid'
                for r_idx, row in enumerate(table_data):
                    for c_idx, cell_val in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = str(cell_val)
                        if r_idx == 0:
                            cell.paragraphs[0].runs[0].font.bold = True

    doc.save(output_path)


def generate_universal_markdown(title: str, content: Dict[str, Any], output_path: str):
    """Generates Markdown file (§5)."""
    md = [f"# {title}\n"]
    overview = content.get("overview") or content.get("summary") or ""
    if overview:
        md.append(f"{overview}\n")

    for sec in content.get("sections") or []:
        sec_title = sec.get("heading") or sec.get("title") or "Section"
        md.append(f"## {sec_title}\n")
        paragraphs = sec.get("paragraphs") or sec.get("content") or []
        if isinstance(paragraphs, str):
            md.append(f"{paragraphs}\n")
        elif isinstance(paragraphs, list):
            for p in paragraphs:
                md.append(f"{p}\n")

        table = sec.get("table")
        if table and len(table) > 1:
            header = table[0]
            md.append("| " + " | ".join(str(h) for h in header) + " |")
            md.append("| " + " | ".join("---" for _ in header) + " |")
            for row in table[1:]:
                md.append("| " + " | ".join(str(c) for c in row) + " |")
            md.append("\n")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md))


def generate_universal_text(title: str, content: Dict[str, Any], output_path: str):
    """Generates Plain Text file (§5)."""
    lines = [f"{title.upper()}\n" + "=" * len(title) + "\n"]
    overview = content.get("overview") or content.get("summary") or ""
    if overview:
        lines.append(f"{overview}\n")

    for sec in content.get("sections") or []:
        sec_title = sec.get("heading") or sec.get("title") or "Section"
        lines.append(f"\n{sec_title.upper()}\n" + "-" * len(sec_title))
        paragraphs = sec.get("paragraphs") or sec.get("content") or []
        if isinstance(paragraphs, str):
            lines.append(f"{paragraphs}")
        elif isinstance(paragraphs, list):
            for p in paragraphs:
                lines.append(f"{p}")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
