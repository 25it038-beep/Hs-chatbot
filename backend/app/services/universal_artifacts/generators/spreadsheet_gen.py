import csv
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_universal_xlsx(title: str, content: Dict[str, Any], output_path: str):
    """Generates a professional XLSX workbook via openpyxl (§6, §34)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:30] or "Data"

    headers = content.get("headers") or []
    rows = content.get("rows") or []

    # If rows is a list of dicts, extract headers
    if rows and isinstance(rows[0], dict) and not headers:
        headers = list(rows[0].keys())
        rows = [[r.get(h, "") for h in headers] for r in rows]

    # Default fallback data if empty
    if not headers and not rows:
        headers = ["ID", "Category", "Description", "Value", "Status"]
        rows = [
            [1, "General", "Initial Metric", 100.0, "Active"],
            [2, "Operations", "System Index", 250.0, "Verified"],
        ]

    # 1. Title Banner
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = title.upper()
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # 2. Header Row
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = str(header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[3].height = 24

    # 3. Data Rows
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    for r_idx, row in enumerate(rows, start=4):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.value = val
            cell.border = thin_border
            if r_idx % 2 == 1:
                cell.fill = alt_fill
            # Align numbers right, text left
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")

    # 4. Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and cell.row > 1:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)


def generate_universal_csv(title: str, content: Dict[str, Any], output_path: str, delimiter: str = ","):
    """Generates standard CSV or TSV file (§6)."""
    headers = content.get("headers") or []
    rows = content.get("rows") or []

    if rows and isinstance(rows[0], dict) and not headers:
        headers = list(rows[0].keys())
        rows = [[r.get(h, "") for h in headers] for r in rows]

    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter=delimiter)
        if headers:
            writer.writerow(headers)
        for row in rows:
            writer.writerow(row)
