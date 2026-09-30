import pytest
import os
import docx
from app.services.nvidia.router import ai_router
from app.services.chat_context.file_adapters.docx_adapter import DOCXAdapter
from app.services.chat_context.context_engine import general_chat_file_context_engine
from app.services.chat_context.contracts import FileProcessingState


def test_router_does_not_hijack_examples_query():
    """Verify queries containing 'examples' or 'give ... examples' are not treated as Wikimedia image search."""
    queries = [
        "give answers for all the question in detailed and with examples",
        "explain operating systems with examples",
        "give me examples of process scheduling algorithms",
        "show me code examples for semaphore in C",
        "give detailed answers for question 1 and 2 with examples",
    ]
    for q in queries:
        decision = ai_router.classify(q)
        assert decision["requires_images"] is False, f"Query '{q}' was falsely flagged as requiring images"
        assert decision["primary_intent"] != "image_search", f"Query '{q}' was falsely classified as image_search"
        task = ai_router.detect_task(q)
        assert task != "web_images", f"Query '{q}' was routed to web_images"


def test_router_prioritizes_files_over_image_search():
    """When files are present in the chat or request, detect_task must never route to web_images."""
    task = ai_router.detect_task("give answers for all the question in detailed and with examples", has_files=True)
    assert task == "chat"


@pytest.mark.asyncio
async def test_docx_adapter_extracts_paragraphs_and_tables(tmp_path):
    """Verify DOCXAdapter extracts questions from paragraphs and tables without losing content."""
    doc_path = str(tmp_path / "OS-Important_questions-SIAT.docx")
    doc = docx.Document()
    doc.add_heading("Operating Systems - Important Questions", level=1)
    doc.add_paragraph("Unit 1: Process Management")
    doc.add_paragraph("1. Explain the difference between process and thread with suitable examples.")
    doc.add_paragraph("2. Describe the Producer-Consumer problem using Semaphores.")

    # Table with questions
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Question No"
    table.rows[0].cells[1].text = "Question Description"
    table.rows[1].cells[0].text = "Q3"
    table.rows[1].cells[1].text = "What is Deadlock? Explain the four necessary conditions."
    table.rows[2].cells[0].text = "Q4"
    table.rows[2].cells[1].text = "Explain Banker's Algorithm with an example allocation matrix."
    doc.save(doc_path)

    adapter = DOCXAdapter()
    result = await adapter.extract(doc_path, "OS-Important_questions-SIAT.docx", "file_os_1")

    assert result.state == FileProcessingState.READY
    assert len(result.chunks) >= 2
    full_content = "\n\n".join(c.content for c in result.chunks)
    assert "difference between process and thread" in full_content
    assert "Deadlock" in full_content
    assert "Banker's Algorithm" in full_content


@pytest.mark.asyncio
async def test_general_chat_file_context_engine_with_docx(tmp_path):
    """End-to-end verification that a docx file attached with 'give answers...' prompt prepares grounded context."""
    doc_path = str(tmp_path / "OS_Questions.docx")
    doc = docx.Document()
    doc.add_heading("Operating System Questions", level=1)
    doc.add_paragraph("1. Define Virtual Memory and Paging.")
    doc.add_paragraph("2. What are the advantages of Multithreading?")
    doc.save(doc_path)

    messages, meta = await general_chat_file_context_engine.prepare_context(
        model_id="llama-3.2-11b",
        user_message="give answers for all the question in detailed and with examples",
        conversation_history=[],
        user_id="user_test_docx",
        chat_id="chat_test_docx",
        attached_files=[{
            "id": "file_os_test",
            "filename": "OS_Questions.docx",
            "storage_path": doc_path,
            "status": "READY",
        }],
        system_prompt="You are HSBot.",
    )

    assert len(messages) >= 2
    # Verify user message and file content are in context
    last_msg = messages[-1]["content"]
    assert "Virtual Memory" in last_msg or "Multithreading" in last_msg
    assert "give answers for all the question in detailed and with examples" in last_msg
    assert meta["chunks_included"] >= 1
